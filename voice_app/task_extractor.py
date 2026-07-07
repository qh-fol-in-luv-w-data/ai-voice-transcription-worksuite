import os
import json
import re
from collections import Counter
from datetime import datetime
from typing import TypedDict, Optional, Annotated
import operator

import requests
from docx import Document
from openai import OpenAI
from langgraph.graph import StateGraph, END
from .constants import get_openai_api_key, get_worksuite_url, get_worksuite_token


# ─────────────────────────────────────────────
# STATE
# ─────────────────────────────────────────────

class AgentState(TypedDict):
    file_path:      str
    api_key:        str
    doc_text:       str
    extracted_data: dict
    session:        object
    frappe_users:   list
    frappe_projects: dict
    created_tasks:  Annotated[list, operator.add]
    errors:         Annotated[list, operator.add]
    tokens_used:    int
    prompt_tokens:  int
    completion_tokens: int


# ─────────────────────────────────────────────
# HELPER: chuẩn hoá & token matching
# ─────────────────────────────────────────────

def _norm(text: str) -> str:
    """Lowercase + bỏ kính ngữ + strip khoảng trắng thừa."""
    if not text:
        return ""
    t = text.strip().lower()
    t = re.sub(r"\b(mr\.?|ms\.?|mrs\.?|anh|chị|chi|ông|bà|ba|em|bạn|chú|cô|bác)\b", "", t)
    return re.sub(r"\s+", " ", t).strip()


def _tokens(text: str) -> set:
    """Tập từ (bỏ từ 1 ký tự)."""
    return {w for w in _norm(text).split() if len(w) > 1}


def _token_overlap(a: str, b: str) -> float:
    """
    Tỉ lệ từ trùng nhau — không phân biệt thứ tự.
    Ví dụ: "Nguyễn Trí Khang" vs "Khang Trí Nguyễn" → 1.0
    """
    ta, tb = _tokens(a), _tokens(b)
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / max(len(ta), len(tb))


def _find_employee(
    name_raw: str,
    designation_hint: Optional[str],
    department_hint: Optional[str],
    employees: list,
    threshold: float = 0.6,
) -> Optional[dict]:
    if not name_raw or not employees:
        return None

    name_clean = _norm(name_raw)
    desig_clean = _norm(designation_hint or "")
    dept_clean  = _norm(department_hint or "")

    scored = []
    for emp in employees:
        name_score = _token_overlap(name_clean, emp.get("employee_name", ""))
        if name_score < threshold:
            continue

        bonus = 0.0
        if desig_clean:
            bonus += 0.3 * _token_overlap(desig_clean, emp.get("designation", ""))
        if dept_clean:
            bonus += 0.2 * _token_overlap(dept_clean,  emp.get("department", ""))

        scored.append((name_score + bonus, emp))

    if not scored:
        return None

    scored.sort(key=lambda x: x[0], reverse=True)

    # Log khi tên trùng — giúp debug
    if len(scored) > 1:
        top, runner = scored[0][0], scored[1][0]
        names = [e.get("employee_name") for _, e in scored[:3]]
        print(f"   🔍 Multi-match '{name_raw}': {names} | scores {[round(s,2) for s,_ in scored[:3]]}")
        if top - runner < 0.1:
            print(f"   ⚠️  Điểm sát ({top:.2f} vs {runner:.2f}) — biên bản nên ghi rõ chức danh/phòng ban!")

    return scored[0][1]


# ─────────────────────────────────────────────
# NODE 1: Đọc file .docx
# ─────────────────────────────────────────────

def node_read_docx(state: AgentState) -> dict:
    print("\n📂 [Node 1] Đọc file:", state["file_path"])
    path = state["file_path"]

    if not os.path.exists(path):
        return {"errors": [f"File không tồn tại: {path}"], "doc_text": ""}

    doc   = Document(path)
    parts = []

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())

    text = "\n".join(parts)
    print(f"   ✅ Đọc được {len(text)} ký tự")
    return {"doc_text": text}


# ─────────────────────────────────────────────
# NODE 2: Trích xuất công việc bằng GPT-4o
# ─────────────────────────────────────────────

def node_extract_tasks(state: AgentState) -> dict:
    model_type = state.get("model_type", "gpt-4o")
    print(f"\n🤖 [Node 2] Phân tích biên bản bằng {model_type}...")

    if not state.get("doc_text"):
        return {"errors": ["Không có nội dung để phân tích"], "extracted_data": {}}

    api_key = get_openai_api_key()
    if not api_key:
        return {"errors": ["Thiếu OPENAI_API_KEY trong config"]}

    client = OpenAI(api_key=api_key)

    prompt = """Bạn là trợ lý phân tích biên bản họp. Đọc nội dung biên bản họp dưới đây và trích xuất tất cả các thông báo và công việc cần xử lý.

Nhiệm vụ của bạn:
- Chỉ trả về các mục công việc và thông báo.
- Bắt buộc phải trích xuất ĐẦY ĐỦ tất cả các Task (nhiệm vụ/công việc) và Noti (thông báo) được nhắc đến trong biên bản, tuyệt đối không được bỏ sót bất kỳ mục nào.
- Mỗi mục là một đầu việc hoặc thông báo riêng, có người thực hiện hoặc người tiếp nhận rõ ràng hoặc là thông báo chung.

Điền thông tin:
- "nguoi_thuc_hien": họ tên ĐẦY ĐỦ chính xác NHƯ TRONG BIÊN BẢN (không rút gọn, không suy diễn, không đảo thứ tự)
- "chuc_danh": chức danh của người đó NẾU được đề cập (ví dụ: "Kế toán trưởng", "Dev", "Lập trình viên"). Null nếu không có.
- "phong_ban": phòng ban / bộ phận NẾU được đề cập. Null nếu không có.
- "ngay_bat_dau": ngày biên bản nếu không ghi cụ thể (dd/mm/yyyy)
- "ngay_ket_thuc": deadline nếu có. Quy đổi: "1-2 buổi"=1 ngày, "3-4 buổi"=2 ngày, "1 tuần"=7 ngày, "vài ngày"=3 ngày. Null nếu không có.
- "note": điểm đặc biệt (rủi ro, điều kiện). Null nếu không có.

Ví dụ: biên bản viết "Anh Minh (Dev)" → nguoi_thuc_hien="Minh", chuc_danh="Dev"
       biên bản viết "Bùi Lễ Văn Minh - Phòng IT" → nguoi_thuc_hien="Bùi Lễ Văn Minh", phong_ban="Phòng IT"

Trả về JSON hợp lệ, KHÔNG có markdown, KHÔNG có giải thích:
{
  "ten_cuoc_hop": "tên cuộc họp ngắn gọn",
  "ngay_hop": "dd/mm/yyyy",
  "items": [
    {
      "id": 1,
      "noi_dung": "mô tả ngắn gọn, rõ ràng",
      "nguoi_thuc_hien": "Họ và tên đầy đủ",
      "chuc_danh": "chức danh hoặc null",
      "phong_ban": "phòng ban hoặc null",
      "ngay_bat_dau": "dd/mm/yyyy",
      "ngay_ket_thuc": "dd/mm/yyyy hoặc null",
      "loai": "task hoặc noti",
      "note": "ghi chú ngắn hoặc null"
    }
  ]
}

Nội dung biên bản họp:
""" + state["doc_text"]

    try:
        response = client.chat.completions.create(
            model=model_type,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.2,
            max_tokens=2000,
        )
        raw = response.choices[0].message.content.strip()

        # Clean JSON
        raw = re.sub(r"```json\s*", "", raw)
        raw = re.sub(r"```", "", raw).strip()
        data = json.loads(raw)

        items_count = len(data.get("items", []))
        print(f"   ✅ Tìm thấy {items_count} mục công việc")
        
        prompt_tokens = response.usage.prompt_tokens if hasattr(response, 'usage') and response.usage else 0
        completion_tokens = response.usage.completion_tokens if hasattr(response, 'usage') and response.usage else 0
        total_tokens = response.usage.total_tokens if hasattr(response, 'usage') and response.usage else 0
        
        return {
            "extracted_data": data,
            "prompt_tokens": prompt_tokens,
            "completion_tokens": completion_tokens,
            "tokens_used": total_tokens
        }
    except Exception as e:
        print(f"   ❌ Lỗi extract tasks: {e}")
        return {"errors": [f"Lỗi extract: {str(e)}"]}


# ─────────────────────────────────────────────
# NODE 3: Login Frappe
# ─────────────────────────────────────────────

def node_login_frappe(state: AgentState) -> dict:
    print("\n🔐 [Node 3] Đăng nhập ERPNext/Frappe...")

    base_url = get_worksuite_url()
    token = get_worksuite_token()

    session = requests.Session()
    session.headers.update({
        "Authorization": f"token {token}",
        "Accept": "application/json"
    })
    print("   ✅ Auth thành công bằng token.")

    return {"session": session}


# ─────────────────────────────────────────────
# NODE 4: Lấy danh sách employees
# ─────────────────────────────────────────────

def node_fetch_users(state: AgentState) -> dict:
    print("\n👥 [Node 4] Lấy danh sách employees...")

    session = state.get("session")
    if not session:
        return {"frappe_users": []}

    base_url = get_worksuite_url()
    resp = session.get(
        f"{base_url}/api/resource/User",
        params={
            "fields":  '["name","full_name","email","enabled"]',
            "filters": '[["enabled","=",1]]',
            "limit":   500,
        },
        timeout=15,
    )

    if resp.status_code != 200:
        print(f"   ⚠️  Không lấy được employees: {resp.status_code}")
        return {"frappe_users": []}

    raw_users = resp.json().get("data", [])
    employees = []
    for u in raw_users:
        employees.append({
            "name": u.get("name"),
            "employee_name": u.get("full_name"),
            "user_id": u.get("email") or u.get("name")
        })
    employees = [e for e in employees if e.get("user_id") and "@" in e.get("user_id", "")]

    # Log tên trùng
    name_counter   = Counter(_norm(e.get("employee_name", "")) for e in employees)
    duplicate_names = {n for n, c in name_counter.items() if c > 1}
    if duplicate_names:
        print(f"   ⚠️  Tên trùng (cần designation/dept để phân biệt): {duplicate_names}")

    print(f"   ✅ {len(employees)} employees có tài khoản")
    return {"frappe_users": employees}


# ─────────────────────────────────────────────
# NODE 4.5: Lấy danh sách Projects
# ─────────────────────────────────────────────

def node_fetch_projects(state: AgentState) -> dict:
    print("\n🏢 [Node 4.5] Lấy danh sách Projects...")
    session = state.get("session")
    if not session:
        return {"frappe_projects": {}}
        
    base_url = get_worksuite_url()
    resp = session.get(
        f"{base_url}/api/resource/Project",
        params={
            "fields": '["name", "project_name", "owner"]',
            "limit": 0,
        },
        timeout=15,
    )
    
    if resp.status_code != 200:
        print(f"   ⚠️  Không lấy được projects: {resp.status_code}")
        return {"frappe_projects": {}}
        
    projects_info = {p["name"]: {"title": p.get("project_name", ""), "owner": p.get("owner")} for p in resp.json().get("data", [])}
    print(f"   ✅ Tìm thấy {len(projects_info)} Projects")
    
    user_projects = {} 
    
    # 1. Gán project cho các owner
    for proj_name, info in projects_info.items():
        owner = info.get("owner")
        title = info.get("title")
        proj_label = title if title else proj_name
        proj_tuple = (proj_label, proj_name)
        
        if owner:
            if owner not in user_projects:
                user_projects[owner] = []
            if proj_tuple not in user_projects[owner]:
                user_projects[owner].append(proj_tuple)
                
    # 2. Gán project từ bảng Project User (chỉ tốn 1 API call)
    resp_users = session.get(
        f"{base_url}/api/resource/Project User",
        params={
            "fields": '["user", "custom_employee", "parent"]',
            "parent": "Project",
            "limit": 0,
        },
        timeout=15,
    )
    
    if resp_users.status_code == 200:
        for row in resp_users.json().get("data", []):
            u_email = row.get("user")
            proj_name = row.get("parent")
            
            if not u_email or not proj_name:
                continue
                
            info = projects_info.get(proj_name, {})
            title = info.get("title")
            proj_label = title if title else proj_name
            proj_tuple = (proj_label, proj_name)
            
            if u_email not in user_projects:
                user_projects[u_email] = []
            if proj_tuple not in user_projects[u_email]:
                user_projects[u_email].append(proj_tuple)
    else:
        print(f"   ⚠️ Lỗi lấy Project User: {resp_users.status_code}")
            
    return {"frappe_projects": user_projects}

# ─────────────────────────────────────────────
# NODE 5: Tạo tasks trên ERPNext (Mới - Tách rời)
# ─────────────────────────────────────────────

def _parse_date(date_str: Optional[str]) -> Optional[str]:
    """dd/mm/yyyy hoặc yyyy-mm-dd → yyyy-mm-dd"""
    if not date_str:
        return None
    # Nếu đã là yyyy-mm-dd thì trả về luôn
    if re.match(r"^\d{4}-\d{2}-\d{2}$", date_str):
        return date_str
    try:
        return datetime.strptime(date_str, "%d/%m/%Y").strftime("%Y-%m-%d")
    except Exception:
        return None


def node_create_tasks(state: AgentState) -> dict:
    print("\n📝 [Node 5] Tạo tasks lên ERPNext...")

    session   = state.get("session")
    data      = state.get("extracted_data", {})
    employees = state.get("frappe_users", [])

    if not session or not data:
        return {"errors": ["Thiếu session hoặc dữ liệu"]}

    items        = data.get("items", [])
    ten_cuoc_hop = data.get("ten_cuoc_hop", "Họp")
    created, errors = [], []

    for item in items:
        # ── Match employee ────────────────────────────────────────────────
        emp = _find_employee(
            name_raw         = item.get("nguoi_thuc_hien", ""),
            designation_hint = item.get("chuc_danh"),
            department_hint  = item.get("phong_ban"),
            employees        = employees,
        )

        assignee_email   = emp.get("user_id")        if emp else None
        assignee_hr_code = emp.get("name")            if emp else None
        assignee_label   = item.get("nguoi_thuc_hien", "N/A")
        emp_desig        = emp.get("designation", "") if emp else ""

        description = f"[{ten_cuoc_hop}]"
        if item.get("note"):
            description += f"\n\n{item['note']}"

        # ── Tạo Task ──────────────────────────────────────────────────────
        payload = {
            "subject":        item["noi_dung"],
            "status":         "Open",
            "description":    description,
            "exp_start_date": _parse_date(item.get("ngay_bat_dau") or data.get("ngay_hop")),
            "exp_end_date":   _parse_date(item.get("ngay_ket_thuc")),
            "custom_assignee":  assignee_hr_code,
            "is_group":       0,
        }
        payload = {k: v for k, v in payload.items() if v is not None and v != ""}

        try:
            base_url = get_worksuite_url()
            resp = session.post(
                f"{base_url}/api/resource/Task",
                json=payload,
                timeout=15,
            )

            if resp.status_code not in (200, 201):
                err = f"'{item['noi_dung'][:40]}': HTTP {resp.status_code} - {resp.text[:300]}"
                print(f"   ❌ {err}")
                errors.append(err)
                continue

            task_name = resp.json().get("data", {}).get("name", "?")

            # # ── Assign via official method (Để gán User và hiển thị Avatar) ──
            assign_ok = True
            if assignee_email and task_name != "?":
                try:
                    session.post(
                        f"{base_url}/api/method/frappe.desk.form.assign_to.add",
                        data={
                            "assign_to": json.dumps([assignee_email]),
                            "doctype": "Task",
                            "name": task_name,
                            "description": "AI Assigned"
                        },
                        timeout=10
                    )
                except:
                    assign_ok = False

            # ── Log ──────────────────────────────────────────────────────
            if emp:
                status = "✓" if assign_ok else "⚠️ assign lỗi"
                match_info = f"{assignee_label} [{emp_desig}] → {assignee_hr_code} ({assignee_email}) {status}"
            else:
                match_info = f"{assignee_label} ⚠️ không tìm thấy employee"

            print(f"   ✅ [{task_name}] '{item['noi_dung'][:50]}' | {match_info}")

            created.append({
                "item_id":        item["id"],
                "task_name":      task_name,
                "title":          item["noi_dung"],
                "assignee_name":  assignee_label,
                "assignee_desig": item.get("chuc_danh"),
                "matched_desig":  emp_desig,
                "hr_emp_code":    assignee_hr_code,
                "assignee_email": assignee_email,
                "assigned":       assign_ok,
                "due_date":       item.get("ngay_ket_thuc"),
            })

        except Exception as e:
            err = f"'{item['noi_dung'][:40]}': {str(e)}"
            print(f"   ❌ {err}")
            errors.append(err)

    print(f"\n   📊 {len(created)}/{len(items)} tasks tạo thành công")
    return {"created_tasks": created, "errors": errors}


# ─────────────────────────────────────────────
# NODE 6: Báo cáo
# ─────────────────────────────────────────────

def node_report(state: AgentState) -> dict:
    print("\n" + "═" * 70)
    data = state.get("extracted_data", {})
    print(f"📄 {data.get('ten_cuoc_hop', 'N/A')}  |  {data.get('ngay_hop', '')}")
    print("═" * 70)

    created = state.get("created_tasks", [])
    errors  = state.get("errors", [])

    if created:
        print(f"\n✅ {len(created)} TASKS ĐÃ TẠO TRÊN ERPNEXT:")
        for t in created:
            due      = f" | Due: {t['due_date']}" if t.get("due_date") else ""
            desig    = f" [{t['matched_desig']}]"  if t.get("matched_desig") else (
                       f" [{t['assignee_desig']}]" if t.get("assignee_desig") else "")
            if t.get("assignee_email"):
                status = "✓" if t.get("assigned") else "⚠️"
                a_str  = f"→ {t['assignee_name']}{desig} | {t['hr_emp_code']} ({t['assignee_email']}) {status}"
            else:
                a_str = f"→ {t['assignee_name']}{desig} ⚠️ chưa assign"
            print(f"   [{t['task_name']}] {t['title'][:50]} {a_str}{due}")

    if errors:
        print(f"\n⚠️  {len(errors)} LỖI:")
        for e in errors:
            print(f"   ❌ {e}")
    return {}


# ─────────────────────────────────────────────
# ROUTERS
# ─────────────────────────────────────────────

def after_read(state: AgentState) -> str:
    return "report" if not state.get("doc_text") else "extract"

def after_extract(state: AgentState) -> str:
    return "report" if not state.get("extracted_data") else "login"

def after_login(state: AgentState) -> str:
    return "report" if not state.get("session") else "fetch_users"


# ─────────────────────────────────────────────
# BUILD & RUN GRAPH
# ─────────────────────────────────────────────

# ─────────────────────────────────────────────
# EXPORT FUNCTIONS
# ─────────────────────────────────────────────

def extract_tasks_only(file_path, model_type="gpt-4o"):
    g = StateGraph(AgentState)
    g.add_node("read_docx",     node_read_docx)
    g.add_node("extract_tasks", node_extract_tasks)
    g.add_node("login",         node_login_frappe)
    g.add_node("fetch_users",   node_fetch_users)
    g.add_node("fetch_projects", node_fetch_projects)

    g.set_entry_point("read_docx")
    g.add_conditional_edges("read_docx",     after_read,    {"extract": "extract_tasks", "report": END})
    g.add_conditional_edges("extract_tasks", after_extract, {"login": "login",           "report": END})
    g.add_conditional_edges("login",         after_login,   {"fetch_users": "fetch_users","report": END})
    g.add_edge("fetch_users", "fetch_projects")
    g.add_edge("fetch_projects", END)

    app = g.compile()
    
    # Trả về kết quả sau khi phân tích
    result = app.invoke({
        "file_path":      file_path,
        "model_type":     model_type,
        "doc_text":       "",
        "extracted_data": {},
        "session":        None,
        "frappe_users":   [],
        "frappe_projects":{},
        "created_tasks":  [],
        "errors":         [],
        "tokens_used":    0,
        "prompt_tokens":  0,
        "completion_tokens": 0,
    })
    
    # Xử lý kết quả để merge thông tin user/project vào extracted_data
    data = result.get("extracted_data", {})
    employees = result.get("frappe_users", [])
    user_projects = result.get("frappe_projects", {})
    
    # Map email to HR_EMP code
    hr_projects_map = {}
    for emp in employees:
        email = emp.get("user_id")
        hr_code = emp.get("name")
        if email and hr_code:
            hr_projects_map[hr_code] = user_projects.get(email, [])
            
    items = []
    for item in data.get("items", []):
        emp = _find_employee(
            name_raw         = item.get("nguoi_thuc_hien", ""),
            designation_hint = item.get("chuc_danh"),
            department_hint  = item.get("phong_ban"),
            employees        = employees,
        )
        
        assignee_email   = emp.get("user_id") if emp else None
        assignee_hr_code = emp.get("name") if emp else None
        assignee_name    = emp.get("employee_name") if emp else item.get("nguoi_thuc_hien", "")
        
        display_assignee = f"{assignee_name} - {assignee_email}" if assignee_email else assignee_name

        item_enriched = {
            "title": item["noi_dung"],
            "assignee_display": display_assignee,
            "start_date": item.get("ngay_bat_dau") or data.get("ngay_hop") or "",
            "end_date": item.get("ngay_ket_thuc") or "",
            "task_type": "noti" if any(k in str(item.get("loai", "")).lower() for k in ["noti", "thông báo"]) else "task",
            "weight": 0,
            "description": item.get('note', '') or ""
        }
        items.append(item_enriched)
        
    usage = {
        "tokens_used": result.get("tokens_used", 0),
        "prompt_tokens": result.get("prompt_tokens", 0),
        "completion_tokens": result.get("completion_tokens", 0)
    }
        
    return items, hr_projects_map, result.get("errors", []), employees, usage


def create_tasks_to_erp(tasks_list):
    """
    Tạo Task từ danh sách đã được user edit trên UI.
    tasks_list là list of dict:
    { "title": str, "assignee_hr_code": str, "assignee_email": str, "project": str, "start_date": str, "end_date": str, "description": str }
    """
    print("\n📝 Bắt đầu tạo Tasks lên ERPNext...")
    
    # Auth bằng token
    base_url = get_worksuite_url()
    token = get_worksuite_token()
    session = requests.Session()
    session.headers.update({
        "Authorization": f"token {token}",
        "Accept": "application/json"
    })

    # Fetch existing tasks to check duplicates
    existing_tasks = []
    try:
        resp_existing = session.get(
            f"{base_url}/api/resource/Task",
            params={
                "fields": '["name", "subject", "project", "custom_assignee"]',
                "filters": '[["status", "in", ["Open", "Working"]]]',
                "limit": 5000,
            },
            timeout=15
        )
        if resp_existing.status_code == 200:
            existing_tasks = resp_existing.json().get("data", [])
    except Exception as e:
        print(f"Lỗi lấy existing tasks: {e}")

    created, errors = [], []
    has_error = False
    
    for item in tasks_list:
        
        # Lấy thẳng project ID vì Dropdown trả về value
        raw_project = item.get("project")
        project_id = raw_project if raw_project else None
            
        # Trích xuất lại assignee_hr_code từ assignee_display nếu user có edit form: "Tên (HR_EMP_...)"
        display_val = item.get("assignee_display", "")
        hr_code = item.get("assignee_hr_code", "")
        if "(HR" in display_val:
            import re
            match = re.search(r"\((HR[-_]EMP[-_][^)]+)\)", display_val)
            if match:
                hr_code = match.group(1)
        elif display_val.startswith("HR_EMP_") or display_val.startswith("HR-EMP-"):
            hr_code = display_val.strip()
            
        subject = item.get("title", "Task không tên")
            
        # Check duplicate
        is_duplicate = False
        for ex in existing_tasks:
            if ex.get("subject") == subject and ex.get("project") == project_id and ex.get("custom_assignee") == hr_code:
                is_duplicate = True
                break
                
        if is_duplicate:
            err = f"'{subject[:40]}': Đã tồn tại (Trùng lặp)"
            print(f"   ⚠️ {err}")
            errors.append(err)
            has_error = True
            break
            
        payload = {
            "subject":        f"[{str(item.get('task_type', 'task')).upper()}] {subject}",
            "status":         "Open",
            "description":    item.get("description", ""),
            "custom_loai_nhiem_vu": item.get("task_type", "task"),
            "exp_start_date": _parse_date(item.get("start_date")),
            "exp_end_date":   _parse_date(item.get("due_date") or item.get("end_date")),
            "custom_assignee":  hr_code,
            "project":        project_id,
            "is_group":       1,
        }
        payload = {k: v for k, v in payload.items() if v is not None and v != ""}

        try:
            resp = session.post(
                f"{base_url}/api/resource/Task",
                json=payload,
                timeout=15,
            )

            if resp.status_code not in (200, 201):
                err_text = resp.text
                if "cannot be later than Project" in err_text or "cannot be before Project" in err_text:
                    # Retry without dates
                    payload.pop("exp_end_date", None)
                    payload.pop("exp_start_date", None)
                    resp = session.post(
                        f"{base_url}/api/resource/Task",
                        json=payload,
                        timeout=15,
                    )
                    
            if resp.status_code not in (200, 201):
                err = f"'{item.get('title', '')[:40]}': HTTP {resp.status_code} - {resp.text[:300]}"
                print(f"   ❌ {err}")
                errors.append(err)
                has_error = True
                break

            task_name = resp.json().get("data", {}).get("name", "?")

            # Gán user
            assign_ok = True
            if item.get("assignee_email") and task_name != "?":
                try:
                    session.post(
                        f"{base_url}/api/method/frappe.desk.form.assign_to.add",
                        data={
                            "assign_to": json.dumps([item["assignee_email"]]),
                            "doctype": "Task",
                            "name": task_name,
                            "description": "AI Assigned"
                        },
                        timeout=10
                    )
                except:
                    assign_ok = False

            created.append({
                "task_name": task_name,
                "title": item["title"],
                "assigned": assign_ok
            })
            print(f"   ✅ Tạo thành công: {task_name}")

        except Exception as e:
            err = f"'{item.get('title', '')[:40]}': {str(e)}"
            print(f"   ❌ {err}")
            errors.append(err)
            has_error = True
            break

    if has_error and created:
        print("\n⏪ Rollback: Xoá các task đã tạo vì có lỗi...")
        for t in created:
            try:
                session.delete(f"{base_url}/api/resource/Task/{t['task_name']}", timeout=10)
                print(f"   Xoá thành công: {t['task_name']}")
            except Exception as e:
                print(f"   Lỗi khi xoá {t['task_name']}: {e}")
        created = []
        errors.append("Đã rollback các task thành công vì có lỗi xảy ra.")

    return {"created_tasks": created, "errors": errors}

def extract_tasks_stateless(docx_path, model_type="gpt-4o-mini"):
    """
    Hàm gọi LangGraph rút gọn (không dính líu Frappe).
    Chỉ dùng LLM để trích xuất Task thô từ nội dung biên bản.
    """
    print(f"\n🚀 Bắt đầu trích xuất Stateless (Không Frappe) với {model_type}...")
    
    # Tạo Graph rút gọn
    from langgraph.graph import StateGraph
    workflow = StateGraph(ExtractionState)
    
    workflow.add_node("read_file", node_read_file)
    workflow.add_node("extract_tasks", node_extract_tasks)
    
    workflow.set_entry_point("read_file")
    workflow.add_edge("read_file", "extract_tasks")
    workflow.set_finish_point("extract_tasks")
    
    app = workflow.compile()
    
    config = {"configurable": {"model_type": model_type}}
    state = ExtractionState(
        file_path=docx_path,
        text_content="",
        raw_llm_output="",
        json_data={},
        employees=[],
        projects={},
        final_tasks=[],
        errors=[]
    )
    
    result = app.invoke(state, config=config)
    
    # Format lại kết quả
    items = []
    data = result.get("json_data", {})
    for item in data.get("items", []):

        items.append({
            "title": item.get("noi_dung", ""),
            "assignee_display": item.get("nguoi_thuc_hien", ""),
            "start_date": item.get("ngay_bat_dau") or data.get("ngay_hop") or "",
            "end_date": item.get("ngay_ket_thuc") or "",
            "task_type": "noti" if any(k in str(item.get("loai", "")).lower() for k in ["noti", "thông báo"]) else "task",
            "weight": 0,
            "description": item.get('note', '') or ""
        })
        
    return items, result.get("errors", [])


def map_speakers_llm(segments: list, speaker_names: list, model_type: str = "gpt-4o") -> dict:
    """
    Dùng LLM map speaker_0, speaker_1... → tên thật dựa vào nội dung hội thoại.
    speaker_names: danh sách tên user nhập trên UI.
    Trả về dict: {"speaker_0": "Mr. Kunalan", "speaker_1": "Ms. Phuong", ...}
    """
    api_key = get_openai_api_key()
    if not api_key or not segments or not speaker_names:
        return {}

    # Lấy danh sách speaker tags có trong transcript
    speaker_tags = sorted(set(s.get("speaker_id", "") for s in segments if s.get("speaker_id")))
    if not speaker_tags:
        return {}

    # Build transcript mẫu (lấy tối đa 60 dòng đầu để không quá dài)
    transcript_lines = []
    for s in segments[:60]:
        spk = s.get("speaker_id", "unknown")
        txt = s.get("text", "").strip()
        if txt:
            transcript_lines.append(f"[{spk}]: {txt}")
    transcript_sample = "\n".join(transcript_lines)

    names_list = ", ".join(speaker_names)
    tags_list  = ", ".join(speaker_tags)

    prompt = f"""Bạn đang phân tích transcript cuộc họp. Hệ thống nhận dạng giọng nói đã tách người nói thành các nhãn: {tags_list}.

Danh sách người tham dự cuộc họp: {names_list}

Dựa vào nội dung hội thoại bên dưới (người được gọi tên, cách xưng hô, ngữ cảnh), hãy xác định mỗi nhãn speaker tương ứng với ai trong danh sách trên.

Quy tắc:
- Chỉ map speaker vào tên có trong danh sách người tham dự
- Nếu không đủ thông tin để xác định chắc chắn, ưu tiên dựa vào thứ tự xuất hiện và ngữ cảnh hội thoại
- Mỗi speaker chỉ map với 1 người, mỗi người chỉ được map 1 lần
- Trả về JSON thuần, không markdown, không giải thích

Transcript (trích):
{transcript_sample}

Trả về JSON:
{{{", ".join(f'"{t}": "tên người"' for t in speaker_tags)}}}"""

    try:
        client = OpenAI(api_key=api_key)
        response = client.chat.completions.create(
            model=model_type,
            messages=[{"role": "user", "content": prompt}],
            temperature=0,
            max_tokens=300,
        )
        raw = response.choices[0].message.content.strip()
        raw = re.sub(r"^```(?:json)?|```$", "", raw, flags=re.MULTILINE).strip()
        mapping = json.loads(raw)
        print(f"[Speaker Map LLM] {mapping}")
        return mapping
    except Exception as e:
        print(f"Lỗi map_speakers_llm: {e}")
        return {}

